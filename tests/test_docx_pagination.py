from docx import Document

from resume_tailor.docx_io import extract_markdown, write_tailored_docx


def test_headings_stay_with_content_across_blank_spacers(tmp_path):
    document = Document()
    document.add_paragraph("Sample Person")
    document.add_paragraph("person@example.invalid")
    document.add_paragraph("PROFESSIONAL EXPERIENCE")
    document.add_paragraph("")
    document.add_paragraph("Example Company | Jan 2020 – Present")
    document.add_paragraph("Software Engineer")
    document.add_paragraph("")
    document.add_paragraph("Built Python services.", "List Bullet")
    document.add_paragraph("\t\t ")
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document.save(source)
    before = source.read_bytes()
    write_tailored_docx(source, extract_markdown(source), output)
    paragraphs = Document(output).paragraphs
    assert all(paragraphs[i].paragraph_format.keep_with_next for i in range(2, 7))
    assert paragraphs[7].paragraph_format.keep_with_next is None
    assert len(paragraphs) == 8
    assert source.read_bytes() == before


def test_manual_page_break_spacers_reflow_without_changing_content_or_styles(tmp_path):
    document = Document()
    document.add_paragraph("Sample Person")
    document.add_paragraph("person@example.invalid")
    document.add_paragraph("TECHNICAL SKILLS", "Heading 1")
    document.add_paragraph("Python, SQL")
    document.add_paragraph("")  # Keep ordinary visual spacing.
    document.add_page_break()
    document.add_paragraph("PROFESSIONAL EXPERIENCE", "Heading 1")
    for company in ("Example Company", "Sample Systems"):
        document.add_paragraph(f"{company} | Jan 2020 – Present")
        document.add_paragraph("Software Engineer")
        bullet = document.add_paragraph("Built reliable Python services for reporting.", "List Bullet")
        bullet.runs[0].italic = True
        if company == "Example Company":
            document.add_page_break()
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document.save(source)
    before = source.read_bytes()
    approved = extract_markdown(source).replace("Built reliable Python services for reporting.",
                                                "Built Python services.")
    write_tailored_docx(source, approved, output)
    written = Document(output)
    assert not written.element.xpath(".//w:br[@w:type='page']")
    assert len([p for p in written.paragraphs if not p.text.strip()]) == 1
    bullets = [p for p in written.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 2
    assert all(p.runs[0].italic for p in bullets)
    assert extract_markdown(output) == approved
    assert source.read_bytes() == before
    headings = [p for p in written.paragraphs if p.style.name == "Heading 1"
                or " | " in p.text or p.text == "Software Engineer"]
    assert all(p.paragraph_format.keep_with_next for p in headings)


def test_page_break_spacer_with_section_properties_is_preserved(tmp_path):
    from docx.oxml import OxmlElement

    document = Document()
    document.add_paragraph("Sample Person")
    document.add_paragraph("person@example.invalid")
    document.add_paragraph("PROFESSIONAL SUMMARY", "Heading 1")
    document.add_paragraph("Python engineer.")
    structural = document.add_page_break()
    structural._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    document.add_paragraph("TECHNICAL SKILLS", "Heading 1")
    document.add_paragraph("Python, SQL")
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document.save(source)
    before = source.read_bytes()
    approved = extract_markdown(source)
    write_tailored_docx(source, approved, output)
    written = Document(output)
    assert len(written.sections) == 2
    assert len(written.element.xpath(".//w:p[w:pPr/w:sectPr]//w:br[@w:type='page']")) == 1
    assert extract_markdown(output) == approved
    assert source.read_bytes() == before
