"""Content omissions must not reach the downloadable Word document."""

import io
from pathlib import Path
from uuid import uuid4

from docx import Document
import pytest

from resume_tailor import server
from resume_tailor.docx_io import extract_markdown
from resume_tailor.pipeline import run_tailoring
from tests.helpers import pack_model_output


@pytest.fixture
def content_run(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "ROOT", tmp_path)
    document = Document()
    for text, style in [
        ("Synthetic Candidate", None),
        ("synthetic@example.invalid", None),
        ("Professional Summary", "Heading 1"),
        ("Built dependable Python data pipelines.", "List Bullet"),
        ("Partnered with analysts to improve data quality.", "List Bullet"),
        ("Technical Skills", "Heading 1"),
        ("Languages: Python, SQL", None),
        ("Cloud: Azure Data Factory, Azure DevOps", None),
        ("Data tools: Kafka, Databricks, Snowflake", None),
        ("Professional Experience", "Heading 1"),
        ("Example Systems | Jan 2021 – Present", None),
        ("Data Engineer", None),
        ("Built ingestion pipelines using Python and SQL.", "List Bullet"),
        ("Monitored ingestion jobs and investigated failures.", "List Bullet"),
        ("Documented data lineage and operational procedures.", "List Bullet"),
        ("Example Analytics | Jan 2017 – Dec 2020", None),
        ("Data Analyst", None),
        ("Developed SQL reports for analysts.", "List Bullet"),
        ("Validated incoming data and resolved quality issues.", "List Bullet"),
        ("Education", "Heading 1"),
        ("B.S. Computing — Example University, 2016", None),
    ]:
        document.add_paragraph(text, style)
    source = tmp_path / "candidate.docx"
    document.save(source)
    original_bytes = source.read_bytes()
    base = extract_markdown(source)
    complete_draft = base.replace(
        "Built dependable Python data pipelines.",
        "Built Python data pipelines with automated validation.",
    ).replace("Cloud: Azure Data Factory, Azure DevOps",
              "Cloud: Azure Data Factory, Azure DevOps, Terraform")

    def run(drafts, *, two_pass=False):
        replies = iter(drafts)
        requests = []

        def complete(messages, **kwargs):
            requests.append(messages)
            return pack_model_output(
                changelog=["Tailored the summary and cloud skills."],
                match="SCORE: 95\ngood: relevant skills", resume=next(replies),
            )

        out_dir = tmp_path / "out" / "runs" / uuid4().hex
        result = run_tailoring(
            job_description="Python SQL data engineering with cloud platforms",
            resume_path=source, out_dir=out_dir, api_key="fake-no-network",
            base_url="https://example.invalid", model="test", complete_fn=complete,
            two_pass=two_pass,
        )
        assert source.read_bytes() == original_bytes
        return result, out_dir, requests

    return complete_draft, run


def _omit_content(draft, omission):
    if omission == "role_bullet":
        return draft.replace("\n- Monitored ingestion jobs and investigated failures.", "")
    if omission == "summary_bullet":
        return draft.replace("\n- Partnered with analysts to improve data quality.", "")
    return draft.replace(", Azure DevOps", "")


def _download(result):
    assert result["ok"], result
    manifest, data = server._read_result(result["run_id"])
    assert manifest["status"] == "ready"
    assert data == Path(result["resume_path"]).read_bytes()
    return Document(io.BytesIO(data))


@pytest.mark.parametrize("omission", ["role_bullet", "summary_bullet", "skill"])
def test_dropped_content_is_restored(content_run, omission):
    # A dropped skill goes straight back; a dropped point first gets the corrective retry.
    draft, run = content_run
    incomplete = _omit_content(draft, omission)
    result, _, requests = run([incomplete, incomplete])

    assert len(requests) == (1 if omission == "skill" else 2)
    assert any(warning.startswith("Restored") for warning in result["warnings"])
    document = _download(result)
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 7
    assert "Monitored ingestion jobs and investigated failures." in bullets
    assert "Partnered with analysts to improve data quality." in bullets
    assert "Built Python data pipelines with automated validation." in bullets
    assert any("Azure DevOps" in p.text and "Terraform" in p.text for p in document.paragraphs)


def test_emptied_role_retries_then_publishes_no_download(content_run):
    draft, run = content_run
    emptied = draft.replace("\n- Developed SQL reports for analysts.", "").replace(
        "\n- Validated incoming data and resolved quality issues.", "")
    result, out_dir, requests = run([emptied, emptied])

    assert len(requests) == 2
    assert "Content was removed" in requests[1][-1]["content"]
    assert not result["ok"]
    assert "Content was removed" in result["error"]
    assert not (out_dir / "result.json").exists()
    assert not list(out_dir.glob("*.docx"))
    with pytest.raises(FileNotFoundError):
        server._read_result(result["run_id"])


def test_polish_cannot_remove_first_pass_added_skill_even_outside_jd(content_run):
    draft, run = content_run
    incomplete_polish = draft.replace(", Terraform", "")
    result, _, requests = run([draft, incomplete_polish], two_pass=True)

    assert len(requests) == 2
    assert any("Terraform" in warning for warning in result["warnings"])
    document = _download(result)
    assert any("Terraform" in p.text for p in document.paragraphs)
    assert "Terraform" not in result["coverage"]["matched"]


def test_two_pass_keeps_the_first_pass_restore_warnings(content_run):
    draft, run = content_run
    result, _, requests = run([draft.replace(", Azure DevOps", ""), draft], two_pass=True)
    assert len(requests) == 2 and not result["polish_error"]
    assert any(warning.startswith("Restored skills") and "Azure DevOps" in warning for warning in result["warnings"])
    assert any("Azure DevOps" in p.text for p in _download(result).paragraphs)
