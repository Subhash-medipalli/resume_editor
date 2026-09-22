import base64
from concurrent.futures import ThreadPoolExecutor
from http.client import HTTPConnection
from http.server import ThreadingHTTPServer
import io
import json
import threading
import time

from docx import Document
import pytest

from resume_tailor import server
from resume_tailor.llm import LLMError
from tests.helpers import pack_model_output


def resume_bytes(name="Sample Candidate"):
    document = Document()
    for text, style in [(name, "Normal"), ("sample@example.invalid", "Normal"),
                        ("PROFESSIONAL SUMMARY:", "Normal"), ("Built Python services.", "List Bullet"),
                        ("Maintained Python services.", "List Bullet"),
                        ("TECHNICAL SKILLS:", "Normal"), ("Tools: Python", "Normal")]:
        document.add_paragraph(text, style)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def with_resume(payload, name="Sample Candidate", filename="sample.docx"):
    body = dict(payload)
    body["resume_name"] = filename
    body["resume_b64"] = base64.b64encode(resume_bytes(name)).decode()
    return body


def reply(messages, **kwargs):
    base = messages[1]["content"].split("## Base resume (source of truth — do not invent beyond this)\n\n", 1)[1]
    return pack_model_output(changelog=["Reworded one bullet"], match="SCORE: 70\npartial: Test reply",
                             resume=base.replace("Built Python services.", "Built reliable Python services."))


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "ROOT", tmp_path)
    monkeypatch.setattr(server, "load_dotenv", lambda *a: None)
    monkeypatch.setattr(server, "complete", reply)
    monkeypatch.setenv("OPENAI_API_KEY", "fake-no-network")
    monkeypatch.chdir(tmp_path)
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    worker = threading.Thread(target=httpd.serve_forever, daemon=True)
    worker.start()

    def raw_request(method, path, payload=None):
        conn = HTTPConnection("127.0.0.1", httpd.server_port, timeout=10)
        try:
            conn.request(method, path, body=json.dumps(payload) if payload is not None else None,
                         headers={"Content-Type": "application/json"})
            res = conn.getresponse()
            body = res.read()
            if res.getheader("Content-Type", "").startswith("application/json"):
                body = json.loads(body)
            return res.status, body
        finally:
            conn.close()
    def request(method, path, payload=None):
        status, body = raw_request(method, path, payload)
        if status == 202:
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline:
                status, progress = raw_request("GET", body["status"])
                if progress["state"] == "complete":
                    return status, progress["result"]
                time.sleep(0.01)
            pytest.fail("run did not finish")
        return status, body
    request.raw = raw_request
    yield request
    httpd.shutdown()
    httpd.server_close()
    worker.join()


def approve(api, result):
    status, approved = api("POST", result["review"], {"reviewed": True, "sha256": result["sha256"]})
    assert status == 200 and approved["ok"]
    return approved["download"]


def test_review_and_downloads_stay_bound_to_their_run(api, monkeypatch):
    status, first = api("POST", "/api/tailor", with_resume({"jd": "Job A"}))
    assert status == 200 and first["ok"] and first["review_required"]
    assert first["download"] is None and first["review_items"]
    first_url = f"/api/download/{first['run_id']}"
    assert api("GET", first_url)[0] == 409
    assert api("POST", first["review"], {"reviewed": True, "sha256": "wrong"})[0] == 409
    assert api("POST", first["review"], {"reviewed": False, "sha256": first["sha256"]})[0] == 409
    assert approve(api, first) == first_url
    status, first_bytes = api("GET", first_url)
    assert status == 200
    _, second = api("POST", "/api/tailor", with_resume({"jd": "Job B"}))
    assert second["run_id"] != first["run_id"]
    assert approve(api, second) != first_url
    assert api("GET", first_url) == (200, first_bytes)
    monkeypatch.setattr(server, "complete", lambda *a, **k: (_ for _ in ()).throw(LLMError("provider failed")))
    status, failed = api("POST", "/api/tailor", with_resume({"jd": "Job C"}))
    assert status == 200 and not failed["ok"] and not failed.get("download")
    assert api("GET", "/api/download")[0] == 404
    assert api("GET", first_url) == (200, first_bytes)


def test_busy_request_is_explicit_and_same_name_uploads_remain_separate(api, monkeypatch):
    started, release = threading.Event(), threading.Event()
    def held_reply(messages, **kwargs):
        if "Candidate One" in messages[1]["content"]:
            started.set()
            assert release.wait(5)
        return reply(messages, **kwargs)
    monkeypatch.setattr(server, "complete", held_reply)
    def run(name):
        return api("POST", "/api/tailor", {"jd": "Python", "resume_name": "resume.docx",
                   "resume_b64": base64.b64encode(resume_bytes(name)).decode()})
    with ThreadPoolExecutor(max_workers=1) as pool:
        first = pool.submit(run, "Candidate One")
        try:
            assert started.wait(5)
            status, busy = run("Candidate Two")
            assert status == 409 and "being processed" in busy["error"]
        finally:
            release.set()
        first_result = first.result()
    second_result = run("Candidate Two")
    for name, (status, result) in zip(("Candidate One", "Candidate Two"), (first_result, second_result)):
        assert status == 200 and result["ok"], result
        status, data = api("GET", approve(api, result))
        assert status == 200
        document = Document(io.BytesIO(data))
        assert document.paragraphs[0].text == name
        assert any("Built reliable" in p.text for p in document.paragraphs)
    assert first_result[1]["run_id"] != second_result[1]["run_id"]


def test_run_snapshot_ignores_later_edits_to_the_saved_base(api, monkeypatch, tmp_path):
    def change_source(messages, **kwargs):
        saved = list((tmp_path / "out/library/resumes").glob("*/*.docx"))
        assert len(saved) == 1
        saved[0].write_bytes(resume_bytes("Updated Candidate"))
        return reply(messages, **kwargs)
    monkeypatch.setattr(server, "complete", change_source)
    _, result = api("POST", "/api/tailor", with_resume({"jd": "Python"}, filename="sample.docx"))
    _, data = api("GET", approve(api, result))
    assert Document(io.BytesIO(data)).paragraphs[0].text == "Sample Candidate"
    saved = next((tmp_path / "out/library/resumes").glob("*/*.docx"))
    assert Document(saved).paragraphs[0].text == "Updated Candidate"
    assert not (tmp_path / "resume").exists()


@pytest.mark.parametrize("failure", ["guardrail", "verification"])
def test_failed_run_has_no_downloadable_artifact(api, monkeypatch, tmp_path, failure):
    if failure == "guardrail":
        def bad_reply(messages, **kwargs):
            return reply(messages, **kwargs) + "\n## EDUCATION\n- Master of Science\n"
        monkeypatch.setattr(server, "complete", bad_reply)
    else:
        monkeypatch.setattr("resume_tailor.docx_io.verify_written_docx", lambda *a: ["synthetic mismatch"])
    status, result = api("POST", "/api/tailor", with_resume({"jd": "Python"}))
    assert status == 200 and not result["ok"]
    assert not result["download"] and not result["review"]
    run_dir = tmp_path / "out/runs" / result["run_id"]
    assert not (run_dir / "result.json").exists()
    assert not list(run_dir.glob("*.docx"))
    assert api("GET", f"/api/download/{result['run_id']}")[0] == 404


def test_changed_result_cannot_be_approved_or_downloaded(api, tmp_path):
    _, result = api("POST", "/api/tailor", with_resume({"jd": "Python"}))
    url = approve(api, result)
    run_dir = tmp_path / "out/runs" / result["run_id"]
    next(run_dir.glob("*.docx")).write_bytes(b"changed after verification")
    assert api("POST", result["review"], {"reviewed": True, "sha256": result["sha256"]})[0] == 409
    assert api("GET", url)[0] == 404


@pytest.mark.parametrize("payload", [["not an object"], {"jd": 123}, {"jd": ""}])
def test_invalid_request_is_rejected_before_allocating_a_run(api, tmp_path, payload):
    assert api("POST", "/api/tailor", payload)[0] == 400
    assert not (tmp_path / "out/runs").exists()


def test_async_status_shows_provider_progress_and_completion(api, monkeypatch):
    waiting, release = threading.Event(), threading.Event()
    def with_progress(messages, **kwargs):
        kwargs["progress"]("Provider unavailable — retrying in 5 seconds")
        waiting.set()
        assert release.wait(5)
        return reply(messages, **kwargs)
    monkeypatch.setattr(server, "complete", with_progress)
    status, started = api.raw("POST", "/api/tailor", with_resume({"jd": "Python"}))
    try:
        assert status == 202
        assert waiting.wait(5)
        status, progress = api.raw("GET", started["status"])
        assert status == 200 and progress["state"] == "working"
        assert "retrying" in progress["message"]
        assert api.raw("GET", "/api/download/" + started["run_id"])[0] == 404
    finally:
        release.set()
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        _, progress = api.raw("GET", started["status"])
        if progress["state"] == "complete":
            break
        time.sleep(0.01)
    assert progress["result"]["ok"] and progress["result"]["review_required"]


def test_oversized_job_description_rejected_before_run_allocation(api, tmp_path):
    status, result = api.raw("POST", "/api/tailor", {"jd": "x" * 60001})
    assert status == 400 and "too long" in result["error"]
    assert not (tmp_path / "out/runs").exists()


def test_health_has_no_built_in_base(api):
    status, body = api("GET", "/api/health")
    assert status == 200 and body["ok"] is True and body["base"] is None


def test_tailor_without_a_base_resume_is_rejected(api, tmp_path):
    status, body = api("POST", "/api/tailor", {"jd": "Python"})
    assert status == 400 and "docx" in body["error"]
    assert not (tmp_path / "out").exists()
    status, body = api("POST", "/api/tailor", {"jd": "Python", "resume_id": "a" * 32})
    assert status == 400 and "Unknown saved resume" in body["error"]
    status, body = api("POST", "/api/tailor", {"jd": "Python", "resume_id": "../secrets"})
    assert status == 400
    assert not (tmp_path / "out").exists()


def test_upload_is_listed_and_can_be_reused(api, tmp_path):
    status, first = api("POST", "/api/tailor", with_resume(
        {"jd": "Python services"}, name="Ada Candidate", filename="Ada Resume.docx",
    ))
    assert status == 200 and first["ok"] and first["resume_id"]
    status, library = api("GET", "/api/library")
    assert status == 200 and len(library["resumes"]) == 1
    saved = library["resumes"][0]
    assert saved["id"] == first["resume_id"]
    assert saved["original_name"] == "Ada Resume.docx"
    assert saved["filename"].endswith(".docx")
    folder = tmp_path / "out/library/resumes" / saved["id"]
    assert (folder / "meta.json").is_file()
    assert list(folder.glob("*.docx"))

    status, second = api("POST", "/api/tailor", {"jd": "Java services", "resume_id": saved["id"]})
    assert status == 200 and second["ok"] and second["resume_id"] == saved["id"]
    assert second["run_id"] != first["run_id"]
    status, library = api("GET", "/api/library")
    assert len(library["resumes"]) == 1

    status, pending = api("GET", "/api/runs")
    by_id = {item["id"]: item for item in pending["runs"]}
    assert by_id[first["run_id"]]["download"] is None
    assert by_id[first["run_id"]]["ok"] is True
    assert "Python services" in by_id[first["run_id"]]["jd_preview"]
    assert by_id[first["run_id"]]["source_name"].endswith(".docx")

    first_url = approve(api, first)
    second_url = approve(api, second)
    status, runs = api("GET", "/api/runs")
    by_id = {item["id"]: item for item in runs["runs"]}
    assert by_id[first["run_id"]]["download"] == first_url
    assert by_id[second["run_id"]]["download"] == second_url
    assert api("GET", by_id[first["run_id"]]["download"])[0] == 200
    assert api("GET", by_id[second["run_id"]]["download"])[0] == 200
    assert Document(io.BytesIO(api("GET", second_url)[1])).paragraphs[0].text == "Ada Candidate"
