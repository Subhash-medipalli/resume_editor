from pathlib import Path
import os

import pytest
from docx import Document

from resume_tailor.cli import _write_outputs, main
from resume_tailor.docx_io import extract_markdown, write_tailored_docx
from resume_tailor.files import atomic_output
from tests.helpers import SAMPLE_RESUME


@pytest.mark.parametrize("alias", ["same", "symlink", "hardlink"])
def test_word_writer_refuses_source_aliases(tmp_path, alias):
    source = tmp_path / "base.docx"
    document = Document()
    document.add_paragraph("Original Candidate")
    document.save(source)
    before = source.read_bytes()
    dest = source if alias == "same" else tmp_path / "output.docx"
    if alias == "symlink":
        dest.symlink_to(source)
    elif alias == "hardlink":
        os.link(source, dest)
    with pytest.raises(ValueError, match="protected source"):
        write_tailored_docx(source, "# Changed Candidate\n", dest)
    assert source.read_bytes() == before


def test_markdown_source_is_checked_before_any_output(tmp_path):
    source = tmp_path / "resume_tailored.md"
    source.write_text(SAMPLE_RESUME)
    with pytest.raises(ValueError, match="protected source"):
        _write_outputs(tailored="# Altered", resume_path=source, out_dir=tmp_path)
    assert source.read_text() == SAMPLE_RESUME


@pytest.mark.parametrize("name", ["CHANGELOG.md", "model.raw.txt", "resume.diff", "resume.rejected.md"])
def test_cli_protects_source_from_diagnostic_outputs(tmp_path, monkeypatch, capsys, name):
    source = tmp_path / name
    source.write_text(SAMPLE_RESUME)
    jd = tmp_path / "jd.txt"
    jd.write_text("Python engineer")
    monkeypatch.setenv("OPENAI_API_KEY", "fake")
    monkeypatch.setattr("resume_tailor.cli.complete", lambda *a, **k: pytest.fail("must reject before the model call"))
    assert main(["--resume", str(source), "--out", str(tmp_path), "--jd", str(jd)]) == 1
    assert "protected source" in capsys.readouterr().err
    assert source.read_text() == SAMPLE_RESUME


def test_writer_does_not_publish_a_failed_verification(tmp_path, monkeypatch):
    source = tmp_path / "base.docx"
    document = Document()
    document.add_paragraph("Original Candidate")
    document.save(source)
    dest = tmp_path / "tailored.docx"
    dest.write_bytes(b"previous result")
    monkeypatch.setattr("resume_tailor.docx_io.verify_written_docx", lambda *a: ["synthetic mismatch"])
    with pytest.raises(ValueError, match="verification failed"):
        write_tailored_docx(source, extract_markdown(source), dest)
    assert dest.read_bytes() == b"previous result"
    assert not list(tmp_path.glob(".resume-tailor-*"))


def test_atomic_write_rechecks_links_at_publication(tmp_path):
    source = tmp_path / "base.md"
    source.write_text("original")
    dest = tmp_path / "out.md"
    with pytest.raises(ValueError, match="protected source"):
        with atomic_output(dest, sources=[source]) as staging:
            staging.write_text("changed")
            dest.symlink_to(source)
    assert source.read_text() == "original"
