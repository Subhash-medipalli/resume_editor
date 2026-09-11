from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from resume_tailor.docx_io import extract_markdown, iter_paragraphs, validate_layout, write_tailored_docx
from resume_tailor.guardrails import apply_guardrails


def test_early_edit_preserves_bold_italic_and_unchanged_run_properties(tmp_path):
    doc = Document()
    doc.add_paragraph("Sample Candidate")
    doc.add_paragraph("Professional Summary", "Heading 1")
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run("Built ")
    bold = paragraph.add_run("Python"); bold.bold = True
    paragraph.add_run(" services for ")
    italic = paragraph.add_run("banking"); italic.italic = True
    paragraph.add_run(".")
    original_styles = [bold._r.rPr.xml, italic._r.rPr.xml, paragraph._p.pPr.xml]
    source, output = tmp_path / "base.docx", tmp_path / "tailored.docx"
    doc.save(source)
    changed = extract_markdown(source).replace("Built ", "Delivered reliable ")
    write_tailored_docx(source, changed, output)
    edited = Document(output).paragraphs[2]
    assert edited.text == "Delivered reliable Python services for banking."
    assert [(r.text, r.bold) for r in edited.runs if r.bold] == [("Python", True)]
    assert [(r.text, r.italic) for r in edited.runs if r.italic] == [("banking", True)]
    assert [edited.runs[1]._r.rPr.xml, edited.runs[3]._r.rPr.xml, edited._p.pPr.xml] == original_styles


def test_nested_tables_extract_in_document_order_but_tailoring_rejects_them(tmp_path):
    doc = Document()
    doc.add_paragraph("Before")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "Cell before"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested"
    table.cell(0, 0).add_paragraph("Cell after")
    doc.add_paragraph("After")
    path = tmp_path / "table.docx"; doc.save(path)
    assert [p.text for p in iter_paragraphs(doc) if p.text] == ["Before", "Cell before", "Nested", "Cell after", "After"]
    with pytest.raises(ValueError, match="tables"):
        validate_layout(path)


@pytest.mark.parametrize("dates", ["2020 – Present", "01/2020 – 12/2023", "January 2020 – Current"])
def test_heading_styles_numbering_and_common_dates_are_preserved(tmp_path, dates):
    doc = Document()
    doc.add_paragraph("Sample Candidate")
    doc.add_paragraph("sample@example.invalid")
    doc.add_paragraph("Professional Experience", "Heading 1")
    doc.add_paragraph("Example Company | " + dates)
    doc.add_paragraph("Software Engineer")
    p = doc.add_paragraph("Built Python services.")
    numbering = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId"); num_id.set(qn("w:val"), "1")
    numbering.append(num_id); p._p.get_or_add_pPr().append(numbering)
    path = tmp_path / "numbered.docx"; doc.save(path)
    validate_layout(path)
    markdown = extract_markdown(path)
    assert "## Professional Experience" in markdown
    assert "### Example Company | " + dates in markdown
    assert "- Built Python services." in markdown
    assert apply_guardrails(markdown, markdown)[1].ok
    assert not apply_guardrails(markdown, markdown.replace(dates, "2015 – 2019"))[1].ok


@pytest.mark.parametrize("layout", ["header", "columns", "job", "plain"])
def test_unsupported_layouts_have_clear_errors(tmp_path, layout):
    doc = Document()
    doc.add_paragraph("Sample Candidate")
    if layout != "plain":
        doc.add_paragraph("Professional Experience", "Heading 1")
    if layout == "header":
        doc.sections[0].header.paragraphs[0].text = "Contact details"
    elif layout == "columns":
        doc.sections[0]._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    elif layout == "job":
        doc.add_paragraph("Company without a date range")
    path = tmp_path / "unsupported.docx"; doc.save(path)
    before = path.read_bytes()
    with pytest.raises(ValueError, match="Unsupported Word layout"):
        validate_layout(path)
    assert path.read_bytes() == before
