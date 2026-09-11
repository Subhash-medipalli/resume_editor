import json
from pathlib import Path

from docx import Document

from resume_tailor.cli import main
from resume_tailor.llm import LLMError
from resume_tailor.pipeline import run_tailoring
from tests.helpers import SAMPLE_RESUME, lightly_tailored, pack_model_output


def run(tmp_path, raw):
    source = tmp_path / "base.md"
    source.write_text(SAMPLE_RESUME)
    return run_tailoring(job_description="Python", resume_path=source,
                         out_dir=tmp_path / "run", api_key="fake", base_url="https://example.invalid",
                         model="test", complete_fn=lambda *a, **k: raw)


def test_summary_describes_final_diff_instead_of_model_claims(tmp_path):
    raw = pack_model_output(changelog=["Invented announcement", "Updated an unrelated qualification", "Raised experience"],
                            match="SCORE: 73\npartial: model estimate", resume=lightly_tailored(SAMPLE_RESUME))
    result = run(tmp_path, raw)
    assert result["ok"] and result["changed_lines"] == 2
    assert len(result["changes"]) == len(result["review_items"]) == 2
    assert all("Edited a passage" in change for change in result["changes"])
    assert "Invented announcement" not in result["changelog"]
    assert result["match_score"] == 73
    assert "Model estimate" in result["score_note"]


def test_score_and_model_assessment_are_withheld_after_reversion(tmp_path):
    changed = lightly_tailored(SAMPLE_RESUME).replace("AWS-hosted platform work.", "AWS-hosted platform work using Terraform.")
    raw = pack_model_output(changelog=["Terraform makes this a great fit"], match="SCORE: 99\ngood: Terraform expert", resume=changed)
    result = run(tmp_path, raw)
    assert result["ok"] and result["changed_lines"] > 0, result
    assert result["match_score"] is None
    assert "withheld" in result["score_note"]
    assert "Terraform expert" not in result["match_line"]
    assert "great fit" not in result["changelog"]
    assert "Terraform" not in Path(result["resume_path"]).read_text()


def test_cli_failure_points_to_its_own_run_and_preserves_prior_result(tmp_path, monkeypatch):
    source, jd, out = tmp_path / "base.md", tmp_path / "jd.txt", tmp_path / "out"
    source.write_text(SAMPLE_RESUME); jd.write_text("Python")
    out.mkdir()
    (out / "resume_tailored.md").write_text("legacy output")
    raw = pack_model_output(changelog=["Edited summary"], match="good: partial alignment", resume=lightly_tailored(SAMPLE_RESUME))
    monkeypatch.setenv("OPENAI_API_KEY", "fake")
    monkeypatch.setattr("resume_tailor.cli.complete", lambda *a, **k: raw)
    args = ["--jd", str(jd), "--resume", str(source), "--out", str(out)]
    assert main(args) == 0
    first = json.loads((out / "latest.json").read_text())
    previous = Path(first["resume"])
    before = previous.read_bytes()
    assert not (out / "resume_tailored.md").exists()
    assert (Path(first["run"]) / "previous-output/resume_tailored.md").read_text() == "legacy output"
    def failure(*a, **k):
        raise LLMError("provider unavailable")
    monkeypatch.setattr("resume_tailor.cli.complete", failure)
    assert main(args) == 1
    latest = json.loads((out / "latest.json").read_text())
    assert latest["status"] == "failed" and latest["resume"] is None
    assert latest["run"] != first["run"]
    assert previous.read_bytes() == before
    assert not (Path(latest["run"]) / "result.json").exists()
    assert "provider unavailable" in (Path(latest["run"]) / "CHANGELOG.md").read_text()
    assert source.read_text() == SAMPLE_RESUME


def test_unsupported_word_layout_is_rejected_before_provider_work(tmp_path):
    document = Document()
    document.add_paragraph("Sample Candidate")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Experience hidden in a table"
    source = tmp_path / "source.docx"; document.save(source)
    def no_model(*a, **k):
        raise AssertionError("must validate the layout first")
    result = run_tailoring(job_description="Python", resume_path=source, out_dir=tmp_path / "out",
                           api_key="fake", base_url="https://example.invalid", model="test", complete_fn=no_model)
    assert not result["ok"] and "Unsupported Word layout" in result["error"]
    assert not list((tmp_path / "out").glob("*.docx"))
