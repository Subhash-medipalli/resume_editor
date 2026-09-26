"""Protected-history regressions through the Word file served for download."""

import io
from pathlib import Path
from uuid import uuid4

from docx import Document
import pytest

from resume_tailor import server
from resume_tailor.docx_io import extract_markdown
from resume_tailor.pipeline import run_tailoring
from tests.helpers import pack_model_output


@pytest.fixture
def tailor(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "ROOT", tmp_path)

    def run(edit, *, certifications=(), polish=None, jd="Python Kubernetes"):
        document = Document()
        for text, style in [
            ("Synthetic Candidate", None),
            ("synthetic@example.invalid", None),
            ("https://www.linkedin.com/in/synthetic-candidate", None),
            ("Professional Summary", "Heading 1"),
            ("Python platform engineer.", None),
            ("Technical Skills", "Heading 1"),
            ("Tools: Python", "List Bullet"),
            ("Professional Experience", "Heading 1"),
            ("Example Systems | Jan 2022 – Present", None),
            ("Data Engineer", None),
            ("Led a 5-person team delivering Python services.", "List Bullet"),
            ("Reduced processing to a 2-hour window.", "List Bullet"),
            ("Built Python pipelines for client Northstar.", "List Bullet"),
            ("Education", "Heading 1"),
            ("B.S. Computing — Example University, 2021", None),
        ]:
            document.add_paragraph(text, style)
        if certifications:
            document.add_paragraph("Certifications", "Heading 1")
            for certification in certifications:
                document.add_paragraph(certification, "List Bullet")
        source = tmp_path / "candidate.docx"
        document.save(source)
        original_bytes = source.read_bytes()
        first = edit(extract_markdown(source))
        drafts = [first] if polish is None else [first, polish(first)]
        replies = iter(
            pack_model_output(changelog=["Tailored to the job description"],
                              match="SCORE: 95\ngood: relevant skills", resume=draft)
            for draft in drafts
        )
        out_dir = tmp_path / "out" / "runs" / uuid4().hex
        result = run_tailoring(
            job_description=jd, resume_path=source, out_dir=out_dir,
            api_key="fake-no-network", base_url="https://example.invalid", model="test",
            complete_fn=lambda *args, **kwargs: next(replies), two_pass=polish is not None,
        )
        assert source.read_bytes() == original_bytes
        return result, out_dir

    return run


def _add_skill(text):
    return text.replace("Tools: Python", "Tools: Python, Kubernetes")


def _download_text(result):
    assert result["ok"], result
    manifest, data = server._read_result(result["run_id"])
    assert manifest["status"] == "ready"
    assert data == Path(result["resume_path"]).read_bytes()
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)


def _assert_repaired(result):
    text = _download_text(result)
    assert "Kubernetes" in text
    assert result["match_score"] is None
    assert "withheld" in result["score_note"]
    assert result["warnings"]
    return text


def _assert_not_downloadable(result, out_dir):
    assert not result["ok"], result
    assert not (out_dir / "result.json").exists()
    assert not list(out_dir.glob("*.docx"))
    with pytest.raises(FileNotFoundError):
        server._read_result(result["run_id"])


@pytest.mark.parametrize("original,changed", [("5-person", "50-person"), ("2-hour", "20-hour")])
def test_download_repairs_hyphenated_numbers_and_keeps_skill_edit(tailor, original, changed):
    result, _ = tailor(lambda base: _add_skill(base).replace(original, changed))
    text = _assert_repaired(result)
    assert original in text and changed not in text


@pytest.mark.parametrize("claim", ["Earned PMP certification.", "PMP-certified."])
def test_download_repairs_new_abbreviated_certifications(tailor, claim):
    result, _ = tailor(lambda base: _add_skill(base).replace("Python platform engineer.", claim))
    text = _assert_repaired(result)
    assert "Python platform engineer." in text and claim not in text


def test_download_cannot_assemble_a_new_certification_from_known_words(tailor):
    claim = "AWS Certified Solutions Architect – Professional with Python expertise."
    result, _ = tailor(
        lambda base: _add_skill(base).replace("Python platform engineer.", claim),
        certifications=["AWS Certified Solutions Architect – Associate", "Project Management Professional (PMP)"],
    )
    text = _assert_repaired(result)
    assert claim not in text
    assert "AWS Certified Solutions Architect – Associate" in text
    assert "Project Management Professional (PMP)" in text


def test_new_degree_in_summary_is_undone_in_the_download(tailor):
    result, _ = tailor(lambda base: _add_skill(base).replace(
        "Python platform engineer.", "PhD-qualified Python platform engineer."
    ))
    text = _assert_repaired(result)
    assert "PhD" not in text and "Python platform engineer." in text


def test_download_restores_separate_contact_url(tailor):
    original = "https://www.linkedin.com/in/synthetic-candidate"
    changed = "https://www.linkedin.com/in/another-candidate"
    result, _ = tailor(lambda base: _add_skill(base).replace(original, changed))
    text = _assert_repaired(result)
    assert original in text and changed not in text


@pytest.mark.parametrize("prefix,suffix", [("", ""), ("- ", ""), ("**", "**"), ("# ", "")])
def test_unmarked_new_job_in_summary_has_no_download(tailor, prefix, suffix):
    result, out_dir = tailor(lambda base: _add_skill(base).replace(
        "Python platform engineer.",
        f"Python platform engineer.\n\n{prefix}Invented Employer | Jan 2022 – Present{suffix}\n**Director of Engineering**",
    ))
    _assert_not_downloadable(result, out_dir)


def test_changed_unpunctuated_client_name_has_no_download(tailor):
    result, out_dir = tailor(lambda base: _add_skill(base).replace("client Northstar", "client Contoso"))
    _assert_not_downloadable(result, out_dir)


def test_download_keeps_first_pass_when_polish_swaps_keywords(tailor):
    result, _ = tailor(
        lambda base: base.replace("Tools: Python", "Tools: Python, Kafka"),
        polish=lambda first: first.replace("Kafka", "RabbitMQ"),
        jd="Required: Python, Kafka, RabbitMQ",
    )
    text = _download_text(result)
    assert "Kafka" in text and "RabbitMQ" not in text
    assert result["polish_error"]
    assert "Kafka" in result["coverage"]["matched"]


def test_download_still_allows_new_jd_skills(tailor):
    result, _ = tailor(lambda base: _add_skill(base).replace(
        "Built Python pipelines", "Built Kubernetes-backed Python pipelines"
    ))
    text = _download_text(result)
    assert "Kubernetes" in text
    assert "Built Kubernetes-backed Python pipelines for client Northstar." in text
    assert result["match_score"] == 95
